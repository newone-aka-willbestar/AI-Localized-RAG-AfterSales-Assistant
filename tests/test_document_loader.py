"""
DocumentLoader 测试。

策略：
- PDF 路径：mock pymupdf4llm 和 fitz，不依赖真实文件
- Word 路径：用 python-docx 在内存中构造真实 .docx，再交给 loader 解析
- 验证：分片数量、元数据字段、格式降级、不支持格式报错
"""
import os
import io
import tempfile
import pytest
from unittest.mock import patch, MagicMock
from pathlib import Path

from src.document_loader import DocumentLoader, SUPPORTED_EXTENSIONS


# ==========================================
# 辅助：创建真实的 .docx 文件
# ==========================================

def make_docx(tmp_path, paragraphs=None, headings=None) -> str:
    """
    在 tmp_path 创建一个简单 .docx 文件，返回路径。
    paragraphs: [{"text": ..., "style": ...}] 列表
    """
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx 未安装，跳过 Word 测试")

    doc = Document()
    if headings:
        for h in headings:
            doc.add_heading(h["text"], level=h.get("level", 1))
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p["text"], style=p.get("style", "Normal"))

    path = str(tmp_path / "test.docx")
    doc.save(path)
    return path


# ==========================================
# SUPPORTED_EXTENSIONS
# ==========================================

class TestSupportedExtensions:

    def test_pdf_supported(self):
        assert ".pdf" in SUPPORTED_EXTENSIONS

    def test_docx_supported(self):
        assert ".docx" in SUPPORTED_EXTENSIONS

    def test_txt_not_supported(self):
        assert ".txt" not in SUPPORTED_EXTENSIONS


# ==========================================
# PDF 路径
# ==========================================

class TestPDFLoading:

    def test_pdf_uses_pymupdf4llm(self, tmp_path):
        """正常情况下走 pymupdf4llm 主路径"""
        loader = DocumentLoader()
        fake_md = "# 第一章\n\n这是正文内容，足够长用于测试分片功能。" * 5

        # 创建空 pdf 占位文件
        pdf_path = str(tmp_path / "test.pdf")
        Path(pdf_path).write_bytes(b"%PDF-1.4 fake")

        with patch("pymupdf4llm.to_markdown", return_value=fake_md):
            docs = loader.load_and_split(pdf_path)

        assert len(docs) > 0
        assert docs[0].metadata["source"] == "test.pdf"
        assert docs[0].metadata["file_type"] == "pdf"

    def test_pdf_falls_back_to_fitz_on_error(self, tmp_path):
        """pymupdf4llm 失败时降级到 fitz"""
        loader = DocumentLoader()
        pdf_path = str(tmp_path / "test.pdf")
        Path(pdf_path).write_bytes(b"%PDF-1.4 fake")

        mock_page = MagicMock()
        mock_page.get_text.return_value = "备用文本内容。" * 20
        mock_doc = MagicMock()
        mock_doc.__iter__ = MagicMock(return_value=iter([mock_page]))
        mock_doc.__enter__ = MagicMock(return_value=mock_doc)
        mock_doc.__exit__ = MagicMock(return_value=False)

        with patch("pymupdf4llm.to_markdown", side_effect=RuntimeError("ONNX 崩溃")):
            with patch("fitz.open", return_value=mock_doc):
                docs = loader.load_and_split(pdf_path)

        assert len(docs) > 0

    def test_pdf_metadata_has_required_fields(self, tmp_path):
        loader = DocumentLoader()
        pdf_path = str(tmp_path / "paper.pdf")
        Path(pdf_path).write_bytes(b"%PDF-1.4 fake")

        with patch("pymupdf4llm.to_markdown", return_value="正文内容。" * 30):
            docs = loader.load_and_split(pdf_path)

        for doc in docs:
            assert "source" in doc.metadata
            assert "chunk_id" in doc.metadata
            assert "file_type" in doc.metadata
            assert "has_table" in doc.metadata


# ==========================================
# Word 路径
# ==========================================

class TestWordLoading:

    def test_docx_loads_paragraphs(self, tmp_path):
        """普通段落被正确提取"""
        path = make_docx(tmp_path, paragraphs=[
            {"text": "这是第一段内容，用于测试 Word 解析功能。" * 5},
            {"text": "这是第二段内容，包含一些论文相关的文字。" * 5},
        ])
        loader = DocumentLoader()
        docs = loader.load_and_split(path)
        assert len(docs) > 0
        full_text = " ".join(d.page_content for d in docs)
        assert "第一段" in full_text or "第二段" in full_text

    def test_docx_metadata_file_type(self, tmp_path):
        """Word 文件的 file_type 元数据为 docx"""
        path = make_docx(tmp_path, paragraphs=[
            {"text": "测试内容。" * 20}
        ])
        loader = DocumentLoader()
        docs = loader.load_and_split(path)
        for doc in docs:
            assert doc.metadata["file_type"] == "docx"
            assert doc.metadata["source"] == "test.docx"

    def test_docx_with_headings(self, tmp_path):
        """标题被转为 Markdown 格式注入"""
        path = make_docx(
            tmp_path,
            headings=[{"text": "研究背景", "level": 1}],
            paragraphs=[{"text": "大语言模型近年来发展迅速。" * 10}],
        )
        loader = DocumentLoader()
        docs = loader.load_and_split(path)
        assert len(docs) > 0

    def test_docx_chunk_ids_sequential(self, tmp_path):
        """分片 ID 从 0 开始连续递增"""
        path = make_docx(tmp_path, paragraphs=[
            {"text": "段落内容。" * 50}
        ])
        loader = DocumentLoader()
        docs = loader.load_and_split(path)
        chunk_ids = [d.metadata["chunk_id"] for d in docs]
        assert chunk_ids == list(range(len(docs)))

    def test_docx_with_table(self, tmp_path):
        """包含表格的 Word 文件不崩溃"""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx 未安装")

        doc = Document()
        doc.add_heading("测试表格", level=1)
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "方法"
        table.cell(0, 1).text = "准确率"
        table.cell(1, 0).text = "BERT"
        table.cell(1, 1).text = "92%"
        table.cell(2, 0).text = "GPT-4"
        table.cell(2, 1).text = "96%"
        path = str(tmp_path / "table_test.docx")
        doc.save(path)

        loader = DocumentLoader()
        docs = loader.load_and_split(path)
        assert len(docs) > 0
        full_text = " ".join(d.page_content for d in docs)
        # 表格内容应被提取
        assert "BERT" in full_text or "GPT" in full_text or "准确率" in full_text


# ==========================================
# 不支持的格式
# ==========================================

class TestUnsupportedFormat:

    def test_txt_raises_value_error(self, tmp_path):
        txt_path = str(tmp_path / "test.txt")
        Path(txt_path).write_text("内容")
        loader = DocumentLoader()
        with pytest.raises(ValueError, match="不支持的文件格式"):
            loader.load_and_split(txt_path)

    def test_xlsx_raises_value_error(self, tmp_path):
        xlsx_path = str(tmp_path / "test.xlsx")
        Path(xlsx_path).write_bytes(b"fake xlsx")
        loader = DocumentLoader()
        with pytest.raises(ValueError, match="不支持的文件格式"):
            loader.load_and_split(xlsx_path)
